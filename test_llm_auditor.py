import os
import docx
import pytest
from app.auditor.auditor import LLMAuditor
from app.auditor.docx_generator import generate_audit_docx

def test_llm_auditor_and_docx_generation(tmp_path):
    """
    Unit test verifying:
    1. Completed investigation state can be passed to auditor.
    2. Auditor returns structured result dict.
    3. DOCX file is generated successfully.
    4. DOCX contains Case ID, Alert ID, Risk Score, Decision, and Audit Conclusion.
    """
    mock_completed_state = {
        "case_id": "CASE_ALT-9999",
        "alert_id": "ALT-9999",
        "entity_id": "CUST-8888",
        "alert_type": "LARGE_AMOUNT",
        "raw_priority_score": 75.0,
        "trigger_evidence": {
            "alert": {"alert_id": "ALT-9999", "customer_id": "CUST-8888", "risk_score": 75.0},
            "customer": {"customer_id": "CUST-8888", "name": "Test User", "risk_level": "HIGH", "occupation": "Software", "account_age_days": 365},
            "transaction": {"transaction_id": "TX-101", "amount": 15000.0, "transaction_type": "WIRE", "status": "COMPLETED"},
            "accounts": [{"account_id": "ACC-01", "status": "ACTIVE"}],
            "beneficiaries": [{"beneficiary_id": "BEN-01", "name": "External Corp"}],
            "devices": [{"device_id": "DEV-01", "device_type": "MOBILE"}]
        },
        "behavioral_metrics": {
            "historical_mean": 2000.0,
            "velocity_z_score": 4.5,
            "velocity_baseline_status": "ELEVATED_VELOCITY",
            "risk_subscores": {"phase1_prior": 75.0, "behavior": 80.0, "graph": 50.0, "kyc": 60.0}
        },
        "graph_metrics": {
            "tx_per_account": 1,
            "beneficiary_dispersion_ratio": 1.0,
            "multi_device_multi_beneficiary_flag": False
        },
        "kyc_notes": "Account age 365 days. Occupation software engineer.",
        "typology_classification": "LARGE_AMOUNT",
        "typology_rationale": "Transaction amount significantly exceeds baseline average.",
        "final_risk_score": 68.75,
        "decision": "REVIEW",
        "dossier": "## SAR Dossier\nCustomer initiated an anomalous $15,000 wire transaction."
    }

    # 1. Instantiate Auditor and run audit
    auditor = LLMAuditor()
    audit_result = auditor.audit_investigation(mock_completed_state)

    # 2. Verify structured audit result fields
    assert audit_result["case_id"] == "CASE_ALT-9999"
    assert audit_result["alert_id"] == "ALT-9999"
    assert "audit_conclusion" in audit_result
    assert audit_result["audit_conclusion"] in ["PASS", "REVIEW REQUIRED", "FAIL"]
    assert isinstance(audit_result["issues_identified"], list)
    assert isinstance(audit_result["auditor_recommendation"], list)

    # 3. Generate DOCX in temp directory
    output_dir = str(tmp_path)
    docx_path = generate_audit_docx(audit_result, mock_completed_state, output_dir=output_dir)

    # 4. Verify file existence
    assert os.path.exists(docx_path), "DOCX report file was not created!"
    assert docx_path.endswith(".docx")

    # 5. Read generated DOCX text content & verify mandatory fields
    doc = docx.Document(docx_path)
    full_text = "\n".join([p.text for p in doc.paragraphs])
    for table in doc.tables:
        for row in table.rows:
            full_text += "\n" + " | ".join([cell.text for cell in row.cells])

    assert "CASE_ALT-9999" in full_text, "DOCX missing Case ID!"
    assert "ALT-9999" in full_text, "DOCX missing Alert ID!"
    assert "68.75" in full_text, "DOCX missing Final Risk Score!"
    assert "REVIEW" in full_text, "DOCX missing Final Decision!"
    assert audit_result["audit_conclusion"] in full_text, "DOCX missing Audit Conclusion!"

    print(f"\n ✓ Test passed! DOCX generated and verified at {docx_path}")
